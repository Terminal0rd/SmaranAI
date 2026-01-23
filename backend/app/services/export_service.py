"""
Export service for PDF and PPT generation
"""
from io import BytesIO
from typing import List, Dict, Any
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import re


class ExportService:
    """Service for exporting sessions to PDF and PPT"""
    
    def __init__(self):
        self.styles = getSampleStyleSheet()
        self._setup_styles()
    
    def _setup_styles(self):
        """Set up custom PDF styles"""
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=24,
            textColor=HexColor('#8b5cf6'),
            spaceAfter=20
        ))
        self.styles.add(ParagraphStyle(
            name='CustomHeading',
            parent=self.styles['Heading2'],
            fontSize=16,
            textColor=HexColor('#6366f1'),
            spaceAfter=10,
            spaceBefore=15
        ))
        self.styles.add(ParagraphStyle(
            name='CustomBody',
            parent=self.styles['Normal'],
            fontSize=11,
            leading=16,
            spaceAfter=8
        ))
        self.styles.add(ParagraphStyle(
            name='Question',
            parent=self.styles['Normal'],
            fontSize=11,
            textColor=HexColor('#3b82f6'),
            spaceBefore=15,
            spaceAfter=5
        ))
    
    def _clean_markdown(self, text: str) -> str:
        """Convert markdown to plain text for PDF"""
        # Remove headers markers but keep text
        text = re.sub(r'^#{1,6}\s*', '', text, flags=re.MULTILINE)
        # Remove bold/italic markers
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        text = re.sub(r'\*(.+?)\*', r'\1', text)
        # Remove code blocks
        text = re.sub(r'```[\s\S]*?```', '', text)
        text = re.sub(r'`(.+?)`', r'\1', text)
        # Remove links but keep text
        text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', text)
        return text.strip()
    
    def generate_pdf(
        self,
        title: str,
        messages: List[Dict[str, Any]],
        concepts: List[str] = None
    ) -> bytes:
        """
        Generate a PDF from session content
        
        Args:
            title: Session title
            messages: List of message dicts
            concepts: List of concepts covered
            
        Returns:
            PDF file as bytes
        """
        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=72
        )
        
        story = []
        
        # Title
        story.append(Paragraph(title, self.styles['CustomTitle']))
        story.append(Spacer(1, 20))
        
        # Concepts section
        if concepts:
            story.append(Paragraph("Key Concepts Covered", self.styles['CustomHeading']))
            concepts_text = " • ".join(concepts)
            story.append(Paragraph(concepts_text, self.styles['CustomBody']))
            story.append(Spacer(1, 20))
        
        # Session content
        story.append(Paragraph("Session Content", self.styles['CustomHeading']))
        story.append(Spacer(1, 10))
        
        for msg in messages:
            if msg['role'] == 'user':
                story.append(Paragraph(
                    f"<b>Question:</b> {self._clean_markdown(msg['content'])}",
                    self.styles['Question']
                ))
            else:
                content = self._clean_markdown(msg['content'])
                # Split into paragraphs
                paragraphs = content.split('\n\n')
                for para in paragraphs:
                    if para.strip():
                        story.append(Paragraph(para, self.styles['CustomBody']))
                story.append(Spacer(1, 15))
        
        doc.build(story)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        
        return pdf_bytes
    
    def generate_pptx(
        self,
        title: str,
        messages: List[Dict[str, Any]],
        concepts: List[str] = None
    ) -> bytes:
        """
        Generate a PowerPoint from session content
        
        Args:
            title: Session title
            messages: List of message dicts
            concepts: List of concepts covered
            
        Returns:
            PPTX file as bytes
        """
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        # Title slide
        title_slide_layout = prs.slide_layouts[6]  # Blank
        slide = prs.slides.add_slide(title_slide_layout)
        
        # Add title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = RGBColor(139, 92, 246)
        p.alignment = PP_ALIGN.CENTER
        
        # Subtitle
        subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(12.333), Inches(0.5))
        tf = subtitle_box.text_frame
        p = tf.paragraphs[0]
        p.text = "Learning Session Export"
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(148, 163, 184)
        p.alignment = PP_ALIGN.CENTER
        
        # Concepts slide
        if concepts:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            
            # Title
            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(0.8))
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = "Key Concepts"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(99, 102, 241)
            
            # Concepts list
            content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.5))
            tf = content_box.text_frame
            tf.word_wrap = True
            
            for i, concept in enumerate(concepts[:12]):
                p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                p.text = f"• {concept}"
                p.font.size = Pt(18)
                p.font.color.rgb = RGBColor(226, 232, 240)
                p.space_after = Pt(8)
        
        # Content slides (one per Q&A pair)
        assistant_messages = [m for m in messages if m['role'] == 'assistant']
        user_messages = [m for m in messages if m['role'] == 'user']
        
        for i, (q, a) in enumerate(zip(user_messages, assistant_messages)):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            
            # Question
            q_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1))
            tf = q_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = f"Q: {self._clean_markdown(q['content'])[:200]}"
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = RGBColor(96, 165, 250)
            
            # Answer
            a_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(12.333), Inches(5))
            tf = a_box.text_frame
            tf.word_wrap = True
            
            content = self._clean_markdown(a['content'])
            # Split into bullet points or paragraphs
            lines = [l.strip() for l in content.split('\n') if l.strip()][:8]
            
            for j, line in enumerate(lines):
                p = tf.add_paragraph() if j > 0 else tf.paragraphs[0]
                # Clean up bullet markers
                if line.startswith('- ') or line.startswith('* '):
                    line = '• ' + line[2:]
                p.text = line[:150]
                p.font.size = Pt(14)
                p.font.color.rgb = RGBColor(226, 232, 240)
                p.space_after = Pt(6)
        
        # Save to bytes
        buffer = BytesIO()
        prs.save(buffer)
        pptx_bytes = buffer.getvalue()
        buffer.close()
        
        return pptx_bytes


# Global instance
export_service = ExportService()
